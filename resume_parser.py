import re
from pathlib import Path
from collections import Counter

import fitz
from docx import Document

try:
    import spacy
    try:
        NLP = spacy.load("en_core_web_sm")
    except Exception:
        NLP = None
except Exception:
    spacy = None
    NLP = None

ALLOWED_EXTENSIONS = {"pdf", "docx"}

SKILLS = [
    "Python", "Java", "C", "C++", "C#", "JavaScript", "TypeScript",
    "HTML", "CSS", "React", "Angular", "Node.js", "Flask", "Django",
    "SQL", "MySQL", "PostgreSQL", "MongoDB", "Oracle", "Power BI",
    "Tableau", "Excel", "Pandas", "NumPy", "Matplotlib", "Seaborn",
    "Scikit-learn", "Machine Learning", "Deep Learning", "NLP",
    "TensorFlow", "PyTorch", "Hugging Face", "AWS", "Azure", "GCP",
    "Git", "GitHub", "Docker", "Kubernetes", "Linux", "REST API",
    "Data Analysis", "Data Science", "Statistics", "Computer Vision",
    "LLM", "Generative AI", "Natural Language Processing",
    "Project Management", "Communication", "Leadership", "Figma", "UI/UX", "User Research", "Prototyping", "Accessibility", "Agile", "Security", "SIEM", "Airflow", "Spark", "Product Management"
]

SECTION_ALIASES = {
    "education": ["education", "academic background", "qualifications"],
    "experience": ["experience", "work experience", "professional experience", "employment"],
    "projects": ["projects", "academic projects", "personal projects"],
    "certifications": ["certifications", "certificates", "licenses"],
    "skills": ["skills", "technical skills", "core skills", "competencies"]
}


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    pages = [page.get_text("text") for page in doc]
    return clean_text("\n".join(pages))


def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    chunks = []

    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))

    return clean_text("\n".join(chunks))


def first_non_empty_line(text):
    for line in text.splitlines():
        line = re.sub(r"[^A-Za-z .'-]", "", line).strip()
        if line and len(line.split()) <= 6:
            if not any(word.lower() in line.lower() for word in [
                "resume", "curriculum vitae", "profile", "email", "phone"
            ]):
                return line
    return None


def extract_email(text):
    match = re.search(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", text)
    return match.group(0) if match else None


def extract_phone(text):
    patterns = [
        r"(?<!\d)(?:\+?\d{1,3}[\s.-]?)?(?:\d[\s.-]?){9,12}\d(?!\d)",
        r"(?<!\d)\d{10}(?!\d)"
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            value = re.sub(r"\s+", " ", match.group(0)).strip()
            digits = re.sub(r"\D", "", value)
            if 10 <= len(digits) <= 13:
                return value
    return None


def extract_skills(text):
    found = []
    lower = text.lower()
    for skill in SKILLS:
        if skill.lower() in lower:
            found.append(skill)
    return sorted(set(found), key=str.lower)


def section_text(text, aliases):
    lines = text.splitlines()
    start = None
    for i, line in enumerate(lines):
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if normalized in aliases or any(normalized == a for a in aliases):
            start = i + 1
            break
    if start is None:
        return ""

    end = len(lines)
    all_headers = set(sum(SECTION_ALIASES.values(), []))
    for i in range(start, len(lines)):
        normalized = re.sub(r"[^a-z ]", "", lines[i].lower()).strip()
        if normalized in all_headers:
            end = i
            break
    return "\n".join(lines[start:end]).strip()


def extract_education(text):
    block = section_text(text, SECTION_ALIASES["education"])
    if not block:
        block = text
    patterns = [
        r"\b(?:B\.?Tech|B\.?E|M\.?Tech|M\.?E|B\.?Sc|M\.?Sc|BCA|MCA|MBA|Ph\.?D|Bachelor(?:'s)?|Master(?:'s)?|Diploma)[^\n]{0,120}"
    ]
    results = []
    for p in patterns:
        results.extend(re.findall(p, block, flags=re.I))
    return list(dict.fromkeys([re.sub(r"\s+", " ", x).strip(" -|") for x in results]))[:8]


def extract_certifications(text):
    block = section_text(text, SECTION_ALIASES["certifications"])
    if not block:
        # Common certification signals when there is no dedicated section.
        matches = re.findall(
            r"(?:certified|certification|certificate)[^.\n]{0,100}",
            text,
            flags=re.I
        )
        return [re.sub(r"\s+", " ", x).strip(" -|:") for x in matches[:10]]

    items = []
    for line in block.splitlines():
        line = line.strip(" •-\t")
        if line:
            items.append(line)
    return items[:10]


def extract_projects(text):
    block = section_text(text, SECTION_ALIASES["projects"])
    items = []
    for line in block.splitlines():
        line = line.strip(" •-\t")
        if line and len(line) > 2:
            items.append(line)
    return items[:12]


def extract_experience(text):
    block = section_text(text, SECTION_ALIASES["experience"])
    source = block or text

    years = re.findall(r"(?i)\b(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)\b", source)
    total_years = max([float(x) for x in years], default=0.0)

    entries = []
    for line in source.splitlines():
        line = line.strip(" •-\t")
        if line and len(line) > 5:
            if re.search(r"(?i)(developer|engineer|analyst|manager|intern|consultant|designer|scientist|lead|associate)", line):
                entries.append(line)

    return {
        "years": total_years,
        "entries": entries[:10]
    }


def extract_name(text):
    # Strong first-line heuristic, then spaCy PERSON entity if available.
    candidate = first_non_empty_line(text)
    if candidate:
        return candidate

    if NLP:
        doc = NLP(text[:10000])
        for ent in doc.ents:
            if ent.label_ == "PERSON":
                return ent.text.strip()
    return None


def extract_candidate_info(text):
    email = extract_email(text)
    phone = extract_phone(text)
    education = extract_education(text)
    skills = extract_skills(text)
    certifications = extract_certifications(text)
    projects = extract_projects(text)
    experience = extract_experience(text)

    if NLP:
        doc = NLP(text[:15000])
        organizations = list(dict.fromkeys(
            ent.text.strip() for ent in doc.ents if ent.label_ == "ORG"
        ))
    else:
        organizations = []

    candidate = {
        "name": extract_name(text),
        "email": email,
        "phone": phone,
        "education": education,
        "skills": skills,
        "experience_years": experience["years"],
        "experience_entries": experience["entries"],
        "certifications": certifications,
        "projects": projects,
        "organizations": organizations,
        "raw_text": text
    }

    # Completeness is a quality indicator, not a claim that every resume is 95% accurate.
    checks = [
        bool(candidate["name"]),
        bool(candidate["email"]),
        bool(candidate["phone"]),
        bool(candidate["education"]),
        bool(candidate["skills"]),
        bool(candidate["certifications"]),
        bool(candidate["projects"]),
        bool(candidate["experience_entries"]) or candidate["experience_years"] > 0
    ]
    candidate["completeness"] = round(sum(checks) / len(checks) * 100, 1)
    return candidate


def process_resume(file_path):
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError("Unsupported file format. Use PDF or DOCX.")

    if not text.strip():
        raise ValueError("No readable text was found in the resume.")

    return extract_candidate_info(text)
